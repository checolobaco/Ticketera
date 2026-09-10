import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

def update_contract_with_ley1493(filepath):
    if not os.path.exists(filepath):
        print(f"File not found: {filepath}")
        return

    doc = docx.Document(filepath)
    
    NAVY = RGBColor(0x0F, 0x17, 0x2A)
    DARK_BLUE = RGBColor(0x1E, 0x3A, 0x8A)

    # Check if clause already exists
    text_full = "\n".join([p.text for p in doc.paragraphs])
    if "1493 DE 2011" in text_full or "1493 de 2011" in text_full:
        print(f"Clause Ley 1493 already exists in: {filepath}")
        return

    # Create new clause heading & paragraph
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(4)
    
    run1 = h.add_run("CLÁUSULA ESPECIAL –")
    run1.bold = True
    run1.font.color.rgb = DARK_BLUE
    run1.font.size = Pt(11.5)
    
    run2 = h.add_run(" NATURALEZA DEL SERVICIO Y EXENCIÓN DE OPERADOR DE BOLETERÍA (LEY 1493 DE 2011):")
    run2.bold = True
    run2.font.color.rgb = NAVY
    run2.font.size = Pt(11.5)

    p = doc.add_paragraph(
        "Las partes acuerdan y reconocen expresamente que EL CONTRATISTA (CloudTickets) es un proveedor de soluciones tecnológicas e infraestructura de software como servicio (SaaS). CloudTickets NO es un operador de boletería, tiquetera pública ni un productor de espectáculos públicos en los términos de la Ley 1493 de 2011 y sus decretos reglamentarios. La plataforma tecnológica es licenciada al CONTRATANTE (Organizador) para que este, de forma autónoma y bajo su propia cuenta, riesgo y responsabilidad, administre la emisión, comercialización, reporte ante el PULEP, retención de tributos y pago de la Contribución Parafiscal sobre las entradas de sus eventos."
    )
    p.paragraph_format.space_after = Pt(8)

    doc.save(filepath)
    print(f"Successfully updated with Ley 1493 clause: {filepath}")

files_to_update = [
    r"c:\0DE\Ticketera\Clients\CONTRATO_SERVICIOS_CLOUDTICKETS_800000_COP.docx",
    r"c:\0DE\Ticketera\CONTRATO_SERVICIOS_CLOUDTICKETS_800000_COP.docx",
    r"c:\0DE\Ticketera\Clients\Harvi_valencia_tulua\CONTRATO_SERVICIOS_CLOUDTICKETS_HARVY_VALENCIA.docx",
    r"c:\0DE\Ticketera\Clients\Marco_Bailey_tulua\BORRADOR CONTRATO_SERVICIOS_CLOUDTICKETS_MARCO_BAILEY.docx"
]

for fp in files_to_update:
    update_contract_with_ley1493(fp)
